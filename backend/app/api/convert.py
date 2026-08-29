import os
import uuid
import hashlib
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from fastapi.responses import StreamingResponse, Response
from sqlalchemy.orm import Session
from typing import List
from datetime import datetime

from ..core.config import settings
from ..models.database import get_db
from ..models.models import User, Conversion, WalletTransaction
from ..parsers import get_parser, reconcile_statement
from ..services.xml_generator import generate_tally_xml
from ..services.excel_exporter import export_transactions_to_excel
from ..schemas.schemas import ConversionResponse, ConversionSummary
from .deps import get_current_user

router = APIRouter(prefix="/convert", tags=["Conversion"])

# 1. Helper to calculate SHA256 file fingerprint
def calculate_file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()

@router.get("/banks", response_model=List[dict])
def get_supported_banks(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    from ..models.models import BankConfig
    configs = db.query(BankConfig).all()
    return [{"bank_key": c.bank_key, "name": c.name, "parser_type": c.parser_type} for c in configs]

@router.post("/upload", response_model=dict)
async def upload_statement(
    bank_name: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # 1. Secure file check (Must be PDF)
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload a PDF file."
        )
        
    content = await file.read()
    
    # Simple validation that it is a valid PDF header
    if not content.startswith(b"%PDF"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is not a valid PDF document."
        )
        
    file_fingerprint = calculate_file_hash(content)
    
    # 2. Duplicate detection check
    existing_conversion = db.query(Conversion).filter(
        Conversion.user_id == current_user.id,
        Conversion.file_fingerprint == file_fingerprint,
        Conversion.status == "SUCCESS"
    ).first()
    
    is_duplicate = existing_conversion is not None
    
    # 3. Save PDF temporarily to scratch/temp_uploads/
    temp_file_id = str(uuid.uuid4())
    temp_filename = f"{temp_file_id}.pdf"
    temp_filepath = os.path.join(settings.TEMP_UPLOAD_DIR, temp_filename)
    
    with open(temp_filepath, "wb") as f:
        f.write(content)
        
    # 4. Parse PDF
    try:
        parser = get_parser(bank_name)
        metadata, transactions = parser.parse(temp_filepath)
    except NotImplementedError as nie:
        # Cleanup temp file
        if os.path.exists(temp_filepath):
            os.remove(temp_filepath)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(nie)
        )
    except Exception as e:
        # Cleanup temp file
        if os.path.exists(temp_filepath):
            os.remove(temp_filepath)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"This PDF does not match the selected {bank_name} statement format. Details: {str(e)}"
        )
        
    # 5. Reconcile
    reconciled = reconcile_statement(metadata.opening_balance, transactions, metadata.closing_balance)
    
    # 6. Format first 15 transactions for UI preview
    preview_txs = []
    for tx in transactions[:15]:
        preview_txs.append({
            "txn_date": tx.txn_date.strftime("%Y-%m-%d"),
            "val_date": tx.val_date.strftime("%Y-%m-%d"),
            "narration": tx.narration,
            "ref_no": tx.ref_no,
            "debit": tx.debit,
            "credit": tx.credit,
            "balance": tx.balance
        })
        
    return {
        "temp_file_id": temp_file_id,
        "file_fingerprint": file_fingerprint,
        "bank_name": metadata.bank_name,
        "account_name": metadata.account_name,
        "account_number": metadata.account_number,
        "branch": metadata.branch,
        "statement_period": metadata.statement_period,
        "pages_count": metadata.pages_count,
        "transactions_count": metadata.transactions_count,
        "opening_balance": metadata.opening_balance,
        "closing_balance": metadata.closing_balance,
        "reconciled": reconciled,
        "is_duplicate": is_duplicate,
        "transactions_preview": preview_txs
    }

@router.post("/execute")
async def execute_conversion(
    temp_file_id: str = Form(...),
    tally_bank_ledger_name: str = Form(...),
    bank_name: str = Form(...),
    counter_ledger_name: str = Form("Suspense"),
    credit_voucher_type: str = Form("Receipt"),
    debit_voucher_type: str = Form("Payment"),
    edited_narrations: str = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if not tally_bank_ledger_name.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Please enter the bank ledger name exactly as it exists in Tally Prime."
        )
        
    temp_filename = f"{temp_file_id}.pdf"
    temp_filepath = os.path.join(settings.TEMP_UPLOAD_DIR, temp_filename)
    
    if not os.path.exists(temp_filepath):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Temporary session expired or file not found. Please upload the PDF again."
        )
        
    # Parse PDF to get page count and transactions
    try:
        parser = get_parser(bank_name)
        metadata, transactions = parser.parse(temp_filepath)
    except Exception as e:
        if os.path.exists(temp_filepath):
            os.remove(temp_filepath)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to parse statement during conversion. Details: {str(e)}"
        )
        
    # Apply user narration edits if provided
    if edited_narrations:
        import json
        try:
            narr_map = json.loads(edited_narrations)
            for idx_str, custom_narr in narr_map.items():
                idx = int(idx_str)
                if 0 <= idx < len(transactions):
                    transactions[idx].narration = custom_narr
        except Exception as e:
            print("Error applying edited narrations:", e)
        
    # 1. Wallet Balance Check
    required_credits = metadata.pages_count
    if current_user.page_credits < required_credits:
        raise HTTPException(
            status_code=status.HTTP_402_PAYMENT_REQUIRED,
            detail=f"Insufficient credits. You need {required_credits} page credits to convert this statement. Your current balance is {current_user.page_credits} pages."
        )
        
    # Calculate file fingerprint to save in log
    with open(temp_filepath, "rb") as f:
        file_fingerprint = calculate_file_hash(f.read())
        
    # 2. Generate XML
    try:
        xml_content = generate_tally_xml(
            transactions, 
            tally_bank_ledger_name,
            counter_ledger_name,
            credit_voucher_type,
            debit_voucher_type
        )
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"XML generation failed: {str(e)}"
        )
    finally:
        # 3. Securely delete PDF immediately after processing
        if os.path.exists(temp_filepath):
            os.remove(temp_filepath)
            
    # 4. Deduct Page Credits (Atomic database updates)
    current_user.page_credits -= required_credits
    
    # Log debit wallet transaction
    wallet_tx = WalletTransaction(
        user_id=current_user.id,
        amount_in_inr=0.0,  # Credit deduction
        pages_purchased=-required_credits,
        type="DEBIT",
        description=f"Converted {metadata.bank_name} statement ({required_credits} pages)"
    )
    db.add(wallet_tx)
    
    # Log conversion history
    conversion = Conversion(
        user_id=current_user.id,
        bank_name=metadata.bank_name,
        account_number=metadata.account_number,
        statement_period=metadata.statement_period,
        pages_count=metadata.pages_count,
        transactions_count=metadata.transactions_count,
        status="SUCCESS",
        file_fingerprint=file_fingerprint
    )
    db.add(conversion)
    db.commit()
    db.refresh(conversion)
    
    # 4.5. Cache the generated XML file locally
    cache_filepath = os.path.join(settings.XML_CACHE_DIR, f"{conversion.id}.xml")
    try:
        with open(cache_filepath, "w", encoding="utf-8") as f:
            f.write(xml_content)
    except Exception as e:
        print("XML caching failed:", e)
        
    # 5. Return Tally XML
    filename = f"Tally_Import_{metadata.bank_name.replace(' ', '_')}_{metadata.statement_period.replace(' ', '_')}.xml"
    
    return Response(
        content=xml_content,
        media_type="application/xml",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.post("/excel")
async def get_excel_export(
    temp_file_id: str = Form(...),
    bank_name: str = Form(...),
    current_user: User = Depends(get_current_user)
):
    temp_filename = f"{temp_file_id}.pdf"
    temp_filepath = os.path.join(settings.TEMP_UPLOAD_DIR, temp_filename)
    
    if not os.path.exists(temp_filepath):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Session expired or file not found. Please upload the PDF again."
        )
        
    try:
        parser = get_parser(bank_name)
        _, transactions = parser.parse(temp_filepath)
        excel_stream = export_transactions_to_excel(transactions)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to export Excel: {str(e)}"
        )
        
    filename = f"Transactions_{bank_name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        excel_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/history", response_model=List[ConversionResponse])
def get_conversion_history(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    history = db.query(Conversion).filter(
        Conversion.user_id == current_user.id
    ).order_by(Conversion.created_at.desc()).all()
    return history

@router.get("/download-xml/{conversion_id}")
def download_cached_xml(
    conversion_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    conversion = db.query(Conversion).filter(
        Conversion.id == conversion_id,
        Conversion.user_id == current_user.id
    ).first()
    if not conversion:
        raise HTTPException(status_code=404, detail="Conversion record not found.")
        
    cache_filepath = os.path.join(settings.XML_CACHE_DIR, f"{conversion.id}.xml")
    if not os.path.exists(cache_filepath):
        raise HTTPException(
            status_code=404, 
            detail="The XML file is no longer available in the local cache. Please re-upload the PDF to convert it again."
        )
        
    with open(cache_filepath, "r", encoding="utf-8") as f:
        xml_content = f.read()
        
    filename = f"Tally_Import_{conversion.bank_name.replace(' ', '_')}_{conversion.statement_period.replace(' ', '_')}.xml"
    
    return Response(
        content=xml_content,
        media_type="application/xml",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/excel-template")
def download_excel_template(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # Gating check: Must be admin or have at least one approved/completed purchase
    has_purchased = current_user.is_admin or db.query(WalletTransaction).filter(
        WalletTransaction.user_id == current_user.id,
        WalletTransaction.type == "PURCHASE",
        WalletTransaction.status.in_(["APPROVED", "COMPLETED"])
    ).first() is not None

    if not has_purchased:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="This free utility is only available to users who have purchased a credit plan. Please buy credits to unlock it."
        )

    import io
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Excel Import Template"
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "Date (DD-MM-YYYY)",
        "Narration",
        "Reference No",
        "Debit (Payment)",
        "Credit (Receipt)",
        "Balance"
    ]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="ffd615", end_color="ffd615", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    for col_idx in range(1, 7):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    sample = ["28-08-2026", "Sample opening receipt transaction", "REF00001", "", "1500.00", "1500.00"]
    ws.append(sample)
    
    for col_idx in range(1, 7):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = Font(name="Calibri", size=11)
        cell.border = thin_border
        if col_idx in [1, 3]:
            cell.alignment = Alignment(horizontal="center")
        elif col_idx in [4, 5, 6]:
            cell.alignment = Alignment(horizontal="right")
            
    column_widths = {"A": 20, "B": 45, "C": 18, "D": 15, "E": 15, "F": 15}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=sheetbuddy_excel_tally_template.xlsx"}
    )

@router.post("/excel-to-xml")
async def convert_excel_to_xml(
    file: UploadFile = File(...),
    bank_ledger_name: str = Form("Bank Account"),
    counter_ledger_name: str = Form("Suspense"),
    credit_voucher_type: str = Form("Receipt"),
    debit_voucher_type: str = Form("Payment"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # Gating check
    has_purchased = current_user.is_admin or db.query(WalletTransaction).filter(
        WalletTransaction.user_id == current_user.id,
        WalletTransaction.type == "PURCHASE",
        WalletTransaction.status.in_(["APPROVED", "COMPLETED"])
    ).first() is not None

    if not has_purchased:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="This free utility is only available to users who have purchased a credit plan. Please buy credits to unlock it."
        )

    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload an Excel file (.xlsx or .xls)."
        )

    content = await file.read()
    
    import io
    import openpyxl
    import datetime
    from ..parsers.base import StandardTransaction
    from ..services.xml_generator import generate_tally_xml

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        
        transactions = []
        for r_idx in range(2, ws.max_row + 1):
            date_val = ws.cell(row=r_idx, column=1).value
            narration_val = ws.cell(row=r_idx, column=2).value
            ref_val = ws.cell(row=r_idx, column=3).value
            debit_val = ws.cell(row=r_idx, column=4).value
            credit_val = ws.cell(row=r_idx, column=5).value
            balance_val = ws.cell(row=r_idx, column=6).value
            
            if date_val is None and narration_val is None and debit_val is None and credit_val is None:
                continue
                
            txn_date = None
            if isinstance(date_val, (datetime.date, datetime.datetime)):
                txn_date = date_val.date() if isinstance(date_val, datetime.datetime) else date_val
            elif isinstance(date_val, str):
                for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
                    try:
                        txn_date = datetime.datetime.strptime(date_val.strip(), fmt).date()
                        break
                    except ValueError:
                        continue
            
            if txn_date is None:
                raise ValueError(f"Invalid date format at row {r_idx}: '{date_val}'. Expected DD-MM-YYYY.")
                
            narration = str(narration_val).strip() if narration_val is not None else ""
            if not narration:
                raise ValueError(f"Missing narration at row {r_idx}.")
                
            def clean_amount(val):
                if val is None or str(val).strip() == "":
                    return None
                try:
                    return float(str(val).replace(",", "").strip())
                except ValueError:
                    raise ValueError(f"Invalid numeric value at row {r_idx}: '{val}'")
                    
            debit = clean_amount(debit_val)
            credit = clean_amount(credit_val)
            balance = clean_amount(balance_val) or 0.0
            
            if debit is None and credit is None:
                raise ValueError(f"Both Debit and Credit columns cannot be empty at row {r_idx}.")
                
            transactions.append(
                StandardTransaction(
                    txn_date=txn_date,
                    val_date=txn_date,
                    narration=narration,
                    ref_no=str(ref_val).strip() if ref_val is not None else None,
                    debit=debit,
                    credit=credit,
                    balance=balance
                )
            )
            
        if not transactions:
            raise ValueError("No valid transactions found in the Excel sheet.")

        xml_data = generate_tally_xml(
            transactions=transactions,
            bank_ledger_name=bank_ledger_name,
            counter_ledger_name=counter_ledger_name,
            credit_voucher_type=credit_voucher_type,
            debit_voucher_type=debit_voucher_type
        )
        
        return Response(
            content=xml_data,
            media_type="application/xml",
            headers={"Content-Disposition": "attachment; filename=excel_tally_import.xml"}
        )

    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Excel parsing failed: {str(e)}"
        )
