import io
from typing import List, Optional, Any
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from ..parsers.base import StandardTransaction, StatementMetadata

def export_transactions_to_excel(
    transactions: List[StandardTransaction],
    metadata: Optional[StatementMetadata] = None,
    bank_name: str = "Bank Statement"
) -> io.BytesIO:
    """
    Exports parsed transactions and an executive reconciliation summary sheet to an Excel workbook.
    """
    wb = Workbook()
    
    # -------------------------------------------------------------
    # SHEET 1: RECONCILIATION SUMMARY
    # -------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Audit & Summary"
    ws_summary.views.sheetView[0].showGridLines = True
    
    font_title = Font(name="Arial", size=14, bold=True, color="000000")
    font_section = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=10, bold=True, color="000000")
    font_regular = Font(name="Arial", size=10, color="000000")
    
    fill_brand = PatternFill(start_color="FFD629", end_color="FFD629", fill_type="solid") # Sheet Buddy Yellow
    fill_black = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    fill_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_green = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    fill_red = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    # Title Block
    ws_summary.merge_cells("A1:D1")
    title_cell = ws_summary["A1"]
    title_cell.value = "SHEET BUDDY — BANK STATEMENT RECONCILIATION AUDIT"
    title_cell.font = font_title
    title_cell.fill = fill_brand
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 32
    
    op_bal = metadata.opening_balance if metadata else (transactions[0].balance if transactions else 0.0)
    cl_bal = metadata.closing_balance if metadata else (transactions[-1].balance if transactions else 0.0)
    tot_dr = sum(t.debit or 0.0 for t in transactions)
    tot_cr = sum(t.credit or 0.0 for t in transactions)
    expected_cl = op_bal + tot_cr - tot_dr
    diff = abs(expected_cl - cl_bal)
    is_rec = (diff < 0.05)
    
    summary_rows = [
        ("Bank Name", metadata.bank_name if metadata else bank_name),
        ("Account Holder", metadata.account_name if metadata else "N/A"),
        ("Account Number", metadata.account_number if metadata else "N/A"),
        ("Statement Period", metadata.statement_period if metadata else "N/A"),
        ("Total Pages Processed", str(metadata.pages_count if metadata else 1)),
        ("Total Transactions", str(len(transactions))),
        ("Opening Balance (A)", f"₹{op_bal:,.2f}"),
        ("Total Credits / Deposits (+B)", f"₹{tot_cr:,.2f}"),
        ("Total Debits / Withdrawals (-C)", f"₹{tot_dr:,.2f}"),
        ("Calculated Closing Balance (A + B - C)", f"₹{expected_cl:,.2f}"),
        ("Statement Closing Balance (D)", f"₹{cl_bal:,.2f}"),
        ("Mathematical Discrepancy", f"₹{diff:,.2f}"),
        ("Audit Reconciliation Status", "RECONCILED (PASS)" if is_rec else "DISCREPANCY DETECTED (CHECK)")
    ]
    
    for r_idx, (k, v) in enumerate(summary_rows, start=3):
        ws_summary.row_dimensions[r_idx].height = 22
        cell_k = ws_summary.cell(row=r_idx, column=1, value=k)
        cell_v = ws_summary.cell(row=r_idx, column=2, value=v)
        cell_k.font = font_bold
        cell_v.font = font_bold if "Status" in k or "Discrepancy" in k else font_regular
        cell_k.border = thin_border
        cell_v.border = thin_border
        cell_k.fill = fill_gray
        
        if k == "Audit Reconciliation Status":
            cell_v.fill = fill_green if is_rec else fill_red
            
    ws_summary.column_dimensions['A'].width = 38
    ws_summary.column_dimensions['B'].width = 45

    # -------------------------------------------------------------
    # SHEET 2: NORMALIZED TRANSACTIONS TABLE
    # -------------------------------------------------------------
    ws_tx = wb.create_sheet(title="Normalized Transactions")
    ws_tx.views.sheetView[0].showGridLines = True
    
    headers = [
        "Txn Date",
        "Value Date",
        "Description / Narration",
        "Reference / UTR",
        "Cheque No",
        "Debit (Payments)",
        "Credit (Receipts)",
        "Balance",
        "Txn Type",
        "Category",
        "Tally Ledger Name",
        "Voucher Type",
        "Confidence",
        "Rec Status",
        "Review Required"
    ]
    ws_tx.append(headers)
    ws_tx.row_dimensions[1].height = 28
    
    # Header format
    for col_idx in range(1, len(headers) + 1):
        c = ws_tx.cell(row=1, column=col_idx)
        c.font = font_section
        c.fill = fill_black
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # Data Rows
    for tx in transactions:
        row = [
            tx.txn_date.strftime("%Y-%m-%d") if tx.txn_date else "",
            tx.val_date.strftime("%Y-%m-%d") if tx.val_date else "",
            tx.narration or "",
            tx.utr or tx.ref_no or "",
            tx.cheque_number or "",
            tx.debit if tx.debit is not None else "",
            tx.credit if tx.credit is not None else "",
            tx.balance,
            tx.transaction_type or ("CREDIT" if tx.credit else "DEBIT"),
            tx.category or "General",
            tx.assigned_ledger or tx.suggested_ledger or "Suspense Account",
            tx.voucher_type or ("Receipt" if tx.credit else "Payment"),
            f"{tx.confidence_score:.0f}%",
            tx.reconciliation_status or "PASS",
            "YES" if tx.requires_review else "NO"
        ]
        ws_tx.append(row)
        r_num = ws_tx.max_row
        ws_tx.row_dimensions[r_num].height = 20
        
        # Format columns
        for c_idx in range(1, len(headers) + 1):
            cell = ws_tx.cell(row=r_num, column=c_idx)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx in [1, 2, 4, 5, 9, 12, 13, 14, 15]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [6, 7, 8]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            # Highlight review rows
            if c_idx == 14 and tx.reconciliation_status == "FAIL":
                cell.fill = fill_red
            elif c_idx == 15 and tx.requires_review:
                cell.fill = fill_brand
                
    # Auto-fit column widths & freeze top header row
    ws_tx.freeze_panes = "A2"
    ws_tx.auto_filter.ref = ws_tx.dimensions
    
    for col in ws_tx.columns:
        col_letter = col[0].column_letter
        max_len = 0
        for cell in col:
            val = str(cell.value or '')
            max_len = max(max_len, len(val))
        ws_tx.column_dimensions[col_letter].width = max(min(max_len + 4, 50), 12)
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
